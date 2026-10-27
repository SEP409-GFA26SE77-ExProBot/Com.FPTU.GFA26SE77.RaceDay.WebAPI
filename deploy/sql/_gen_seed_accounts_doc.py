from pathlib import Path

from docx import Document
from docx.shared import Pt

out = Path(__file__).resolve().parents[1] / "SEED_ACCOUNTS.docx"
doc = Document()

style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)

doc.add_heading("RaceDay — danh sách account seed", level=1)
p = doc.add_paragraph()
p.add_run("Nguồn: ").bold = True
p.add_run("deploy/clear-and-seed.ps1  →  deploy/sql/clear_and_seed.sql")
doc.add_paragraph("Password chung tất cả account: MatKhau@123")
doc.add_paragraph("Email đã verify sẵn — login không cần OTP.")

rows = [
    ("admin", "MatKhau@123", "admin@raceday.vn", "Admin"),
    ("cuong.le", "MatKhau@123", "cuong.le@raceday.vn", "Organizer"),
    ("an.nguyen", "MatKhau@123", "an.nguyen@raceday.vn", "User"),
    ("binh.tran", "MatKhau@123", "binh.tran@raceday.vn", "User"),
    ("dung.pham", "MatKhau@123", "dung.pham@raceday.vn", "User"),
    ("hoa.le", "MatKhau@123", "hoa.le@raceday.vn", "User"),
    ("khanh.vo", "MatKhau@123", "khanh.vo@raceday.vn", "User"),
    ("lan.pham", "MatKhau@123", "lan.pham@raceday.vn", "User"),
    ("minh.do", "MatKhau@123", "minh.do@raceday.vn", "User"),
    ("nga.bui", "MatKhau@123", "nga.bui@raceday.vn", "User"),
    ("phuc.hoang", "MatKhau@123", "phuc.hoang@raceday.vn", "User"),
    ("quynh.dao", "MatKhau@123", "quynh.dao@raceday.vn", "User"),
    ("son.ngo", "MatKhau@123", "son.ngo@raceday.vn", "User"),
    ("tram.ly", "MatKhau@123", "tram.ly@raceday.vn", "User"),
    ("uyen.mai", "MatKhau@123", "uyen.mai@raceday.vn", "User"),
]

table = doc.add_table(rows=1 + len(rows), cols=4)
table.style = "Table Grid"
for i, h in enumerate(["Username", "Password", "Gmail / Email", "Role"]):
    cell = table.rows[0].cells[i]
    cell.text = h
    for run in cell.paragraphs[0].runs:
        run.bold = True

for r_i, row in enumerate(rows, start=1):
    for c_i, val in enumerate(row):
        table.rows[r_i].cells[c_i].text = val

note = doc.add_paragraph()
note.add_run("\nGhi chú").bold = True
doc.add_paragraph(
    "• 11 User đầu (an.nguyen → son.ngo): đã ĐK RaceDay Grand Marathon (Confirmed + Paid + BIB)."
)
doc.add_paragraph("• tram.ly, uyen.mai: chỉ có hồ sơ — dùng test đăng ký mới.")
doc.add_paragraph(
    "• Organizer cuong.le: sở hữu Grand Marathon (RegistrationOpen) + Trail Draft."
)

doc.save(out)
print(f"Wrote {out} ({out.stat().st_size} bytes)")
