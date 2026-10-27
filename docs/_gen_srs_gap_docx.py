# -*- coding: utf-8 -*-
"""Generate docs/SRS_3.2_3.8.2_Gap_Context.docx from gap matrix."""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = r"d:\SEP\SEP_WebAPI\docs\SRS_3.2_3.8.2_Gap_Context.docx"

doc = Document()
for section in doc.sections:
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2)
    section.right_margin = Cm(2)


def set_run_font(run, size=11, bold=False, color=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color


def add_heading_custom(text, level=1):
    p = doc.add_paragraph()
    run = p.add_run(text)
    sizes = {0: 18, 1: 14, 2: 12}
    color = RGBColor(0xC2, 0x41, 0x0C) if level <= 1 else RGBColor(0x1F, 0x29, 0x37)
    set_run_font(run, size=sizes.get(level, 12), bold=True, color=color)
    p.paragraph_format.space_before = Pt(14 if level else 0)
    p.paragraph_format.space_after = Pt(6)


def add_body(text, italic=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, size=11)
    run.italic = italic
    p.paragraph_format.space_after = Pt(6)


def shade_cell(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_text(cell, text, bold=False, size=9, white=False):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text)
    set_run_font(
        run,
        size=size,
        bold=bold,
        color=RGBColor(255, 255, 255) if white else None,
    )
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)


def add_table(headers, rows, header_fill="EA580C"):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, bold=True, size=9, white=True)
        shade_cell(table.rows[0].cells[i], header_fill)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            set_cell_text(table.rows[ri + 1].cells[ci], val, size=9)
            if ri % 2 == 1:
                shade_cell(table.rows[ri + 1].cells[ci], "FFF7ED")
            if len(headers) >= 5 and ci == 4 and str(val).strip().lower() == "high":
                shade_cell(table.rows[ri + 1].cells[ci], "FEE2E2")
    doc.add_paragraph()


add_heading_custom("SRS §3.2–3.8.2 — Context cần sửa (BE / FE / Docs)", 0)
add_body("Nguồn: Report3_Software Requirement Specification.docx")
add_body("Phạm vi đọc: 3.2 → hết 3.8.2")
add_body(
    "Đối chiếu: RaceDay modular monolith (SEP_WebAPI), FE User / Admin, "
    "OCR (SEP_Service_Photo) theo .cursor/rules + AGENTS.md."
)
add_body(
    "Chỉ liệt kê gap / lệch nghiệp vụ hoặc SRS lỗi. Mục đã khớp "
    "(register/login/forgot, event wizard cơ bản, PayOS, check-in volunteer, "
    "photo OCR/publish/search/dispute…) không đưa vào bảng.",
    italic=True,
)

add_heading_custom("A. Lỗi tài liệu SRS (sửa Report3 trước)", 1)
add_table(
    ["Mục SRS", "Vấn đề", "Hướng sửa Docs"],
    [
        [
            "3.2.4 Profile",
            "Body copy nhầm thành Logout (trùng 3.2.5)",
            "Viết lại Profile: xem/sửa hồ sơ, medical/profile fields theo BR-17–25",
        ],
        [
            "3.2.5 Logout",
            "Trùng nội dung với 3.2.4 hiện tại",
            "Giữ flow logout; bỏ duplicate",
        ],
        [
            "3.3.4 Configure Event BIB",
            "Trigger/description paste nhầm Admin approval",
            "Viết đúng: Organizer cấu hình BIB range theo distance",
        ],
        [
            "3.3.5 Distance Map",
            "Cùng paste Admin approval",
            "Viết đúng: upload/cấu hình map lộ trình",
        ],
        [
            "3.3.4 (× nhiều)",
            "Staff / Licenses / Content / Supplies / Approval đều đánh số 3.3.4",
            "Đánh số lại 3.3.6–3.3.10",
        ],
        [
            "3.7.6 Photo Detail & Dispute",
            "Copy nguyên 3.7.5 Search by BIB",
            "Viết: mở chi tiết, tải/share, dispute ảnh sai, photographer xử lý",
        ],
        [
            "3.8.2 Configure Platform Audit",
            "Title = Audit; body = moderation nội dung",
            "Tách: (a) Audit log read, (b) Content/report moderation — hoặc đổi title",
        ],
    ],
)

add_heading_custom("B. Gap sản phẩm vs nghiệp vụ SRS", 1)
add_table(
    ["Mục", "Context", "SRS kỳ vọng", "Hiện trạng", "Sev", "Layer"],
    [
        [
            "3.2.2",
            "Login identifier",
            "Email + password",
            "FE/BE login bằng username (email chỉ lúc register)",
            "Med",
            "BE/FE/Docs",
        ],
        [
            "3.2.5",
            "Logout token",
            "Session/token invalid",
            "FE clear cookie/JWT; không revoke server-side",
            "Low",
            "BE",
        ],
        [
            "3.3.x",
            "Notify sau approve/reject event",
            "Thông báo Organizer khi có hỗ trợ notify",
            "Chủ yếu Audit/integration event; thiếu email/in-app",
            "Med",
            "BE",
        ],
        [
            "3.4.1",
            "Slot hold khi chờ thanh toán",
            "Giữ chỗ trong cửa sổ PayOS (nếu support)",
            "Chỉ check capacity lúc create; không hold",
            "Med",
            "BE",
        ],
        [
            "3.4.3",
            "BIB + QR sau payment verified",
            "Flow: Payment → confirm → BIB → QR",
            "Payment → pending_approval; BIB/QR sau Organizer approve registration",
            "High",
            "BE/FE/Docs",
        ],
        [
            "3.4.4",
            "Ticket khi eligible",
            "QR/BIB khi đủ điều kiện",
            "Khớp product nếu giữ bước approve; lệch SRS nếu hiểu sau payment là có vé",
            "Med",
            "FE/Docs",
        ],
        [
            "3.5.3",
            "Record race finish",
            "Staff/Organizer ghi finish (UI và/hoặc API)",
            "BE có POST .../finish; FE không wire",
            "High",
            "FE",
        ],
        [
            "3.5.3",
            "Finish overwrite",
            "Không ghi đè ngoài ý muốn; có rule sửa",
            "Repeat → AlreadyScanned; không API sửa finish",
            "Med",
            "BE",
        ],
        [
            "3.6.1",
            "Import/manage results",
            "Organizer quản lý / CSV import",
            "Không import API; Results FE phần lớn prototype",
            "High",
            "BE/FE",
        ],
        [
            "3.6.2",
            "Publish results",
            "Trạng thái unpublished vs published",
            "Public leaderboard = mọi Finished; không gate publish",
            "High",
            "BE/FE",
        ],
        [
            "3.6.x",
            "Result complaint",
            "Khiếu nại kết quả (BR-73/75)",
            "FE dispute route chưa BE; /disputes = photo dispute",
            "Med",
            "BE/FE",
        ],
        [
            "3.7.4",
            "Unpublish gallery",
            "Publish và unpublish",
            "Domain có Unpublish; thiếu API/FE",
            "Med",
            "BE/FE",
        ],
        [
            "3.7.5",
            "Empty vs OCR đang chạy",
            "Không nói không có ảnh chỉ vì OCR chưa xong",
            "Search chỉ published; thiếu signal đang xử lý",
            "Low",
            "FE",
        ],
        [
            "3.8.1",
            "Admin user mgmt",
            "CRUD/role/reset theo quyền",
            "BE đủ hơn; Admin FE chủ yếu list + lock/unlock",
            "Med",
            "FE",
        ],
        [
            "3.8.2",
            "Audit vs moderation",
            "Body SRS = duyệt nội dung báo cáo",
            "Admin FE Audit = read-only log; không moderation queue",
            "High",
            "BE/FE/Docs",
        ],
        [
            "NFR/Check-in",
            "QR chống giả",
            "Ticket không đoán được",
            "QR JSON plain; chưa ký",
            "Med",
            "BE",
        ],
        [
            "Admin extras",
            "Stats / reconciliation",
            "Vận hành admin",
            "FE demo/static",
            "Med",
            "FE",
        ],
    ],
)

add_heading_custom("C. Ưu tiên backlog (đề xuất)", 1)
for item in [
    "Chốt nghiệp vụ 3.4.3: giữ Organizer approve rồi mới BIB/QR → sửa SRS; hoặc đổi BE sang auto-BIB sau payment verified như SRS.",
    "Results publish gate (3.6.2) + organizer results UI/API (3.6.1).",
    "Wire FE race finish (3.5.3).",
    "Sửa Report3 docs (3.2.4, 3.3 numbering, 3.7.6, 3.8.2).",
    "Photo unpublish API/FE; Admin user actions + moderation nếu giữ wording 3.8.2.",
]:
    p = doc.add_paragraph(style="List Number")
    p.clear()
    run = p.add_run(item)
    set_run_font(run, size=11)

add_heading_custom("D. Đã khớp đáng kể (không cần sửa gấp)", 1)
for item in [
    "3.2.1 Register, verify email, hash password, default role",
    "3.3 Create event + profile distances/BIB/map/docs/content/supplies + Admin approve event",
    "3.4 PayOS checkout + payment status tách registration",
    "3.5.1–3.5.2 Check-in QR / manual (Admin volunteer + BE)",
    "3.7.1–3.7.5 Upload → OCR → review/correct → publish → search by BIB (+ download/dispute FE gần đây)",
]:
    p = doc.add_paragraph(style="List Bullet")
    p.clear()
    run = p.add_run(item)
    set_run_font(run, size=11)

add_heading_custom("E. Ghi chú", 1)
add_body(
    "Gap High quan trọng nhất cần chốt product: §3.4.3 — giữ bước Organizer approve "
    "(cập nhật SRS) hay auto-BIB sau payment (đổi BE)."
)
add_body(
    "Tài liệu này chỉ mang tính đối chiếu gap; không thay thế Report3 SRS gốc."
)

doc.save(OUT)
print("Wrote", OUT)
