"""Generate docs/API_Supplies_Map_Staff.docx — run: python docs/generate_supplies_map_api_docx.py"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

OUT = Path(__file__).resolve().parent / "API_Supplies_Map_Staff.docx"


def h(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def p(doc: Document, text: str, *, bold: bool = False) -> None:
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.bold = bold
    run.font.size = Pt(11)


def code(doc: Document, text: str) -> None:
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(9)


def bullet(doc: Document, text: str) -> None:
    doc.add_paragraph(text, style="List Bullet")


def main() -> None:
    doc = Document()
    title = doc.add_heading("RaceDay API — Vật phẩm tiếp tế & Map course (staff)", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p(doc, "Phiên bản tài liệu: 2026-10-25 · Host Docker :5080 · Modular monolith Event + Map + IAM")
    p(doc, "Nguồn sự thật: SEP_WebAPI (Event supplies + Map EnsureCanRead) khớp FE SEP_Pull_Dev_FE_User.")

    h(doc, "0. Tóm tắt kiến trúc")
    bullet(doc, "Vật phẩm tiếp tế nằm trong module Event (không module mới) — scope per-event.")
    bullet(doc, "Map course: MỘT API GET hiện có; mở rộng quyền đọc cho Active staff (không tách API theo role).")
    bullet(doc, "Login unverified: message → «Vui lòng xác thực.» (không «nhập OTP» trên form login).")
    bullet(doc, "OTP: chỉ audit trong mục 4 — không đổi code trong đợt này.")

    # --- SUPPLIES ---
    h(doc, "1. Event — Vật phẩm tiếp tế")
    p(doc, "Auth chung: Bearer JWT · Role Organizer (chủ giải) hoặc Admin · Policy PublicUser trên group organizer.")
    p(doc, "Base: /api/v1/event/organizer/{eventId}/supplies")

    h(doc, "1.1 GET — danh sách vật phẩm", 2)
    p(doc, "Roles: Organizer (owner của eventId) | Admin", bold=True)
    p(doc, "Request: không body. Path param eventId (uuid).")
    p(doc, "Response 200 — JSON array:")
    code(
        doc,
        """[
  {
    "id": "a1b2c3d4-e5f6-7890-abcd-ef1234567890",
    "name": "Nước khoáng thiên nhiên 500ml",
    "quantity": 80,
    "unit": "thùng",
    "notes": "Thùng 24 chai. Phân bổ đồng đều cho các trạm tiếp nước."
  }
]""",
    )
    bullet(doc, "id (string, Guid): khóa bản ghi DB — FE khai báo string.")
    bullet(doc, "name (string): tên vật phẩm.")
    bullet(doc, "quantity (int ≥ 0): số thùng.")
    bullet(doc, "unit (string): mặc định «thùng».")
    bullet(doc, "notes (string|null): quy cách / ghi chú đóng gói.")
    p(doc, "Status:")
    bullet(doc, "200 OK — danh sách (có thể []).")
    bullet(doc, "401 — thiếu/invalid token.")
    bullet(doc, "400/DomainException — chưa login / không phải chủ giải / không tìm thấy giải.")

    h(doc, "1.2 PUT — thay thế toàn bộ danh sách", 2)
    p(doc, "Roles: Organizer (owner) | Admin", bold=True)
    p(doc, "Request JSON:")
    code(
        doc,
        """{
  "supplies": [
    {
      "id": "sup-1",
      "name": "Nước khoáng thiên nhiên 500ml",
      "quantity": 80,
      "unit": "thùng",
      "notes": "Thùng 24 chai. Phân bổ đồng đều cho các trạm tiếp nước."
    },
    {
      "id": null,
      "name": "Gel năng lượng chạy bộ thể thao",
      "quantity": 12,
      "unit": "thùng",
      "notes": "Thùng 24 gói. Kiểm tra hạn sử dụng trước khi giao nhận."
    }
  ]
}""",
    )
    bullet(doc, "supplies (array, required): tối đa 50 phần tử.")
    bullet(doc, "id (string|null, optional): nếu parse được Guid hợp lệ → giữ Id; FE «sup-…» → BE sinh Guid mới.")
    bullet(doc, "name (required, ≤200), quantity (≥0), unit (≤32, default thùng), notes (≤1000).")
    p(doc, "Hành vi: xóa hết supplies của event rồi insert lại (replace-all).")
    p(doc, "Response 200 — mảng giống GET (id luôn Guid string sau khi lưu).")
    p(doc, "Status: 200 | 401 | 400 (validation / không phải organizer / >50 items).")

    h(doc, "1.3 POST — thêm 1 vật phẩm", 2)
    p(doc, "Roles: Organizer (owner) | Admin", bold=True)
    p(doc, "Request JSON:")
    code(
        doc,
        """{
  "id": null,
  "name": "Ly giấy sinh học dùng 1 lần",
  "quantity": 20,
  "unit": "thùng",
  "notes": "Thùng 1.000 ly giấy tự phân hủy."
}""",
    )
    bullet(doc, "Cùng validation từng field như PUT; từ chối nếu đã đủ 50 item.")
    p(doc, "Response 200 — một object EventSupplyItemDto (id Guid mới nếu client id không phải Guid).")
    p(doc, "Status: 200 | 401 | 400.")

    h(doc, "1.4 DELETE — xóa 1 vật phẩm", 2)
    p(doc, "Roles: Organizer (owner) | Admin", bold=True)
    p(doc, "Path: .../supplies/{supplyId} với supplyId = Guid.")
    p(doc, "Response: 204 No Content.")
    p(doc, "Status: 204 | 401 | 400 (không tìm thấy / không thuộc event).")

    h(doc, "1.5 Seed DB (test)", 2)
    p(
        doc,
        "EventDataSeeder.SeedDemoSuppliesAsync: mọi race_events chưa có supplies nhận 5 dòng khớp FE DEFAULT_SUPPLIES "
        "(80/50/30/12/6 thùng). Idempotent.",
    )

    # --- MAP ---
    h(doc, "2. Map — Course")
    p(doc, "Base: /api/v1/map/events/{eventId}/distances/{distanceId}/course")
    p(doc, "FE compat (cùng handler): /api/v1/event/organizer/{eventId}/distances/{distanceId}/course")

    h(doc, "2.1 GET course", 2)
    p(
        doc,
        "Roles / điều kiện đọc (MỘT API — không tách): Admin | Organizer owner | "
        "Active Event staff (Coordinator / Volunteer / Photographer) kể cả draft | "
        "mọi caller khi status public (Approved|RegistrationOpen|Closed|Completed). "
        "GET map không bắt buộc JWT; staff draft cần Bearer có UserId liên kết staff Active.",
        bold=True,
    )
    p(doc, "Request: path eventId, distanceId — không body.")
    p(doc, "Response 200 — CourseConfigDto:")
    code(
        doc,
        """{
  "title": "Fun Run 5K",
  "distanceId": "22222222-bbbb-4bbb-8bbb-bbbbbbbbbbb1",
  "distanceKm": 5,
  "waypoints": [ { "lat": 13.78, "lng": 109.21 } ],
  "polyline": [ { "lat": 13.78, "lng": 109.21 }, { "lat": 13.79, "lng": 109.22 } ],
  "checkpoints": [
    {
      "id": "cp-start",
      "name": "Start",
      "km": 0,
      "lat": 13.78,
      "lng": 109.21,
      "type": "start",
      "cutoff": null
    }
  ],
  "updatedAt": "2026-10-25T12:00:00Z"
}""",
    )
    bullet(doc, "title: tên course / cự ly.")
    bullet(doc, "distanceId: Guid cự ly (có thể null khi empty stub).")
    bullet(doc, "distanceKm: độ dài km.")
    bullet(doc, "waypoints / polyline: mảng {lat,lng}; chưa có course → [].")
    bullet(doc, "checkpoints[].type: start|cp|aid|finish|water|medical|cutoff.")
    bullet(doc, "updatedAt: UTC lần cập nhật entity (hoặc now nếu empty).")
    p(doc, "Status: 200 | 400 (không tìm thấy giải/cự ly hoặc không đủ quyền đọc draft).")

    h(doc, "2.2 PUT course (không đổi RBAC đợt này)", 2)
    p(doc, "Roles: Organizer owner | Admin (RequireAuthorization PublicUser + handler check).", bold=True)
    p(doc, "Request UpsertCourseRequest:")
    code(
        doc,
        """{
  "title": "Fun Run 5K",
  "distanceId": "22222222-bbbb-4bbb-8bbb-bbbbbbbbbbb1",
  "distanceKm": 5,
  "waypoints": [ { "lat": 13.78, "lng": 109.21 } ],
  "polyline": [ { "lat": 13.78, "lng": 109.21 } ],
  "checkpoints": [
    { "id": "aid-1", "name": "Trạm tiếp nước", "km": 2.5, "lat": 13.785, "lng": 109.215, "type": "aid" }
  ]
}""",
    )
    bullet(doc, "distanceKm: 0–500; waypoints/polyline/checkpoints ≤ 500 phần tử.")
    bullet(doc, "checkpoint.type ∈ {start,cp,aid,finish,water,medical,cutoff}.")
    p(doc, "Response 200 — CourseConfigDto (có updatedAt sau upsert).")
    p(doc, "Status: 200 | 401 | 400 (validation / không phải organizer).")

    # --- LOGIN ---
    h(doc, "3. Login message (IAM)")
    p(doc, "Khi User.Status=Pending hoặc !EmailVerified, EnsureCanLogin ném:")
    code(doc, '"Tài khoản chưa xác thực email. Vui lòng xác thực."')
    p(
        doc,
        "Context: form login không phải nơi nhập OTP — user cần verify-email / resend OTP. "
        "Đổi «nhập OTP» → «xác thực» là đúng UX.",
    )

    # --- OTP AUDIT ---
    h(doc, "4. OTP audit (check-only — không sửa code)")
    p(doc, "Nguồn: EmailOtp.DefaultTtlMinutes=10, MaxAttempts=5; EmailOtpService RNG 6 số; hash SHA256(email|purpose|code).")
    bullet(doc, "Độ dài 6 + crypto RNG — đạt.")
    bullet(doc, "TTL 10 phút (khuyến nghị 5–15’) — đạt.")
    bullet(doc, "Giới hạn 5 lần verify rồi invalidate — đạt.")
    bullet(doc, "So sánh FixedTimeEquals — đạt.")
    bullet(doc, "Thiếu: resend cooldown / rate-limit OTP riêng (chỉ public 200/min).")
    bullet(doc, "Thiếu: pepper/HMAC server secret (hash thuần SHA256).")
    bullet(doc, "Gap nhỏ: resend-otp có thể lộ email không tồn tại.")
    p(doc, "Kết luận: đủ quy chuẩn OTP cơ bản cho MVP; production nên bổ sung throttle + pepper.")

    h(doc, "5. Docker / test nhanh")
    bullet(doc, "docker compose --profile api up -d --build api (deploy/).")
    bullet(doc, "Seed: SELECT * FROM event.event_supplies; kỳ vọng 5 dòng/giải.")
    bullet(doc, "FE: /organizer/events/{id}/profile/supplies — GET/PUT thật, không cần localStorage fallback.")

    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
